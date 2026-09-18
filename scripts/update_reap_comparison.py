#!/usr/bin/env python3
"""Refresh the REAP comparison table from completed local evaluation artifacts."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
EVAL_ROOT = ROOT / "artifacts" / "eval"
RESULT_ROOT = EVAL_ROOT / "fgcluster_reap_localfr_gl_0p5"
CSV_PATH = EVAL_ROOT / "qwen3_reap_expert_merge_comparison.csv"
MD_PATH = EVAL_ROOT / "qwen3_reap_expert_merge_comparison.md"
XLSX_PATH = EVAL_ROOT / "qwen3_reap_expert_merge_comparison.xlsx"
LEGACY_TARGET_VARIANT = "REAP-conditioned local-FR cluster + GL 50%"
TARGET_VARIANT = "Geodesic-Cluster-GL Merge"


def latest(paths: list[Path]) -> Path | None:
    return max(paths, key=lambda path: path.stat().st_mtime) if paths else None


def read_evalscope_score(dataset: str) -> float | None:
    report = latest(list(RESULT_ROOT.glob(f"evalscope_results/*/reports/*/{dataset}.json")))
    if report is None:
        return None
    return 100.0 * float(json.loads(report.read_text())["score"])


def read_livecode_score() -> float | None:
    result = RESULT_ROOT / "Scenario.codegeneration_1_0.2_eval.json"
    if not result.exists():
        return None
    payload = json.loads(result.read_text())
    return 100.0 * float(payload[0]["pass@1"])


def fmt(value: float | None) -> str:
    return "" if value is None else f"{value:.2f}"


def update_row(row: dict[str, str]) -> None:
    livecode = read_livecode_score()
    gsm8k = read_evalscope_score("gsm8k")
    math500 = read_evalscope_score("math_500")

    row["LiveCode"] = fmt(livecode)
    row["GSM8K"] = fmt(gsm8k)
    row["MATH-500"] = fmt(math500)

    evalplus = float(row["Eval+ Avg"])
    row["Paper Code Avg"] = fmt((evalplus + livecode) / 2 if livecode is not None else None)
    row["Math Avg"] = fmt((gsm8k + math500) / 2 if gsm8k is not None and math500 is not None else None)

    pending: list[str] = []
    if math500 is None:
        pending.append("MATH-500 running")
    if not list(RESULT_ROOT.glob("wildbench*/**/stats.json")):
        pending.append("WildBench pending")
    row["Status"] = "; ".join(pending) if pending else "Evaluation complete"


def write_csv(fieldnames: list[str], rows: list[dict[str, str]]) -> None:
    with CSV_PATH.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def write_markdown(fieldnames: list[str], rows: list[dict[str, str]]) -> None:
    intro = (
        "# Qwen3 REAP Pruning and Expert-Merging Comparison\n\n"
        "Paper REAP is pruning, while Geodesic-Cluster-GL is an expert-merging method. "
        "Percent scores. Paper Code Avg = mean(Eval+, LiveCode); custom Overall excludes "
        "LiveCode/WildBench/math.\n\n"
    )
    header = "| " + " | ".join(fieldnames) + " |\n"
    separator = "| " + " | ".join("---" for _ in fieldnames) + " |\n"
    body = "".join(
        "| " + " | ".join(row.get(field, "") or "—" for field in fieldnames) + " |\n"
        for row in rows
    )
    MD_PATH.write_text(intro + header + separator + body)


def update_xlsx(fieldnames: list[str], rows: list[dict[str, str]]) -> None:
    workbook = load_workbook(XLSX_PATH)
    sheet = workbook["Scores"]
    headers = {str(cell.value): cell.column for cell in sheet[1]}
    target_row = next(
        index for index in range(2, sheet.max_row + 1)
        if sheet.cell(index, headers["Variant"]).value in {TARGET_VARIANT, LEGACY_TARGET_VARIANT}
    )
    row = next(item for item in rows if item["Variant"] == TARGET_VARIANT)
    numeric = set(fieldnames) - {"Variant", "Technique", "Source", "Artifact", "Status"}
    for field in fieldnames:
        value: Any = row[field]
        if field in numeric and value != "":
            value = float(value)
        elif value == "":
            value = None
        sheet.cell(target_row, headers[field]).value = value

    summary_name = "Paper comparison"
    if summary_name in workbook.sheetnames:
        del workbook[summary_name]
    summary = workbook.create_sheet(summary_name, 0)

    columns = [
        "Model", "HumanEval", "HumanEval+", "MBPP", "MBPP+", "Eval+ Avg",
        "MC Avg", "LiveCodeBench", "Paper Code Avg", "WildBench", "GSM8K",
        "MATH-500", "Math Avg",
    ]
    metric_fields = [
        "HumanEval", "HumanEval+", "MBPP", "MBPP+", "Eval+ Avg", "MC Avg",
        "LiveCode", "Paper Code Avg", "WildBench", "GSM8K", "MATH-500", "Math Avg",
    ]
    variants = [
        "Dense parent (paper)",
        "REAP pruning 50% (paper)",
        TARGET_VARIANT,
    ]
    display_names = {
        "Dense parent (paper)": "Dense paper baseline",
        "REAP pruning 50% (paper)": "REAP pruning, paper",
        TARGET_VARIANT: TARGET_VARIANT,
    }
    by_variant = {row["Variant"]: row for row in rows}

    summary.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    title = summary.cell(1, 1, "Qwen3 — Paper REAP vs. Geodesic-Cluster-GL Merge")
    title.font = Font(name="Aptos Display", size=20, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="172554")
    title.alignment = Alignment(horizontal="left", vertical="center")
    summary.row_dimensions[1].height = 34

    summary.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columns))
    subtitle = summary.cell(
        2,
        1,
        "Scores are percentages • Paper Code Avg = mean(Eval+ Avg, LiveCodeBench) • "
        "WildBench judge protocols differ",
    )
    subtitle.font = Font(name="Aptos", size=10, italic=True, color="475569")
    subtitle.alignment = Alignment(horizontal="left")
    summary.row_dimensions[2].height = 24

    header_fill = PatternFill("solid", fgColor="2563EB")
    for column, label in enumerate(columns, 1):
        cell = summary.cell(4, column, label)
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    summary.row_dimensions[4].height = 32

    thin = Side(style="thin", color="CBD5E1")
    for row_index, variant in enumerate(variants, 5):
        source = by_variant[variant]
        values: list[Any] = [display_names[variant]] + [source[key] for key in metric_fields]
        row_fill = "EFF6FF" if variant == TARGET_VARIANT else ("F8FAFC" if row_index % 2 else "FFFFFF")
        for column, value in enumerate(values, 1):
            cell = summary.cell(row_index, column)
            if column > 1 and value not in ("", None):
                cell.value = float(value)
                cell.number_format = "0.00"
            elif column > 1:
                cell.value = "Pending" if variant == TARGET_VARIANT else "—"
            else:
                cell.value = value
            cell.fill = PatternFill("solid", fgColor=row_fill)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(
                horizontal="left" if column == 1 else "center",
                vertical="center",
                wrap_text=True,
            )
            cell.font = Font(
                name="Aptos",
                size=10,
                bold=variant == TARGET_VARIANT,
                color="0F172A",
            )
        summary.row_dimensions[row_index].height = 28

    summary.merge_cells(start_row=9, start_column=1, end_row=9, end_column=len(columns))
    target = by_variant[TARGET_VARIANT]
    completed = [
        f"LiveCodeBench complete ({target['LiveCode']})" if target["LiveCode"] else "LiveCodeBench pending",
        f"GSM8K complete ({target['GSM8K']})" if target["GSM8K"] else "GSM8K pending",
        f"MATH-500 complete ({target['MATH-500']})" if target["MATH-500"] else "MATH-500 pending",
    ]
    note = summary.cell(
        9,
        1,
        "Latest status: " + "; ".join(completed) +
        "; WildBench unavailable with the current Gemma runtime.",
    )
    note.fill = PatternFill("solid", fgColor="FEF3C7")
    note.font = Font(name="Aptos", size=10, italic=True, color="92400E")
    note.alignment = Alignment(wrap_text=True, vertical="center")
    summary.row_dimensions[9].height = 32

    summary.freeze_panes = "B5"
    summary.auto_filter.ref = f"A4:{get_column_letter(len(columns))}7"
    summary.sheet_view.showGridLines = False
    summary.column_dimensions["A"].width = 31
    for column in range(2, len(columns) + 1):
        summary.column_dimensions[get_column_letter(column)].width = 15
    summary.column_dimensions["H"].width = 18
    summary.column_dimensions["I"].width = 18
    workbook.save(XLSX_PATH)


def main() -> None:
    with CSV_PATH.open(newline="") as handle:
        reader = csv.DictReader(handle)
        fieldnames = list(reader.fieldnames or [])
        rows = list(reader)

    target = next(
        row for row in rows
        if row["Variant"] in {TARGET_VARIANT, LEGACY_TARGET_VARIANT}
    )
    target["Variant"] = TARGET_VARIANT
    update_row(target)
    write_csv(fieldnames, rows)
    write_markdown(fieldnames, rows)
    update_xlsx(fieldnames, rows)

    print(
        f"Updated {TARGET_VARIANT}: LiveCode={target['LiveCode'] or 'pending'}, "
        f"GSM8K={target['GSM8K'] or 'pending'}, MATH-500={target['MATH-500'] or 'pending'}, "
        f"WildBench=pending"
    )


if __name__ == "__main__":
    main()
